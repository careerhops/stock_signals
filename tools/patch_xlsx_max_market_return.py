from __future__ import annotations

import shutil
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook


INPUT_XLSX = Path("/Users/madhubhatt/Desktop/Weekly Buy Tracker and Technical Analysis.xlsx")
OHLCV_CSV = Path("/Users/madhubhatt/Desktop/NSE_daily_ohlcv (1).csv")
OUTPUT_DIR = Path("/Users/madhubhatt/Documents/stock_signals/outputs/max_market_return")
OUTPUT_XLSX = OUTPUT_DIR / "Weekly Buy Tracker and Technical Analysis - Max Market Return.xlsx"

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"main": MAIN_NS, "rel": REL_NS, "pkg": PKG_REL_NS}

ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)


def qname(tag: str) -> str:
    return f"{{{MAIN_NS}}}{tag}"


def parse_workbook_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def parse_number(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_symbol(value):
    return str(value).strip().upper() if value is not None else ""


def excel_serial(dt: date) -> int:
    return (dt - date(1899, 12, 30)).days


def column_index(column_letters: str) -> int:
    total = 0
    for char in column_letters:
        total = total * 26 + ord(char.upper()) - ord("A") + 1
    return total


def cell_column(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha())
    return column_index(letters)


def make_number_cell(ref: str, value, style: str | None = None) -> ET.Element:
    cell = ET.Element(qname("c"), {"r": ref})
    if style is not None:
        cell.set("s", style)
    v = ET.SubElement(cell, qname("v"))
    if isinstance(value, float):
        v.text = f"{value:.12g}"
    else:
        v.text = str(value)
    return cell


def make_inline_string_cell(ref: str, value: str, style: str | None = None) -> ET.Element:
    cell = ET.Element(qname("c"), {"r": ref, "t": "inlineStr"})
    if style is not None:
        cell.set("s", style)
    inline = ET.SubElement(cell, qname("is"))
    text = ET.SubElement(inline, qname("t"))
    text.text = value
    return cell


def set_cell(row: ET.Element, cell: ET.Element):
    ref = cell.attrib["r"]
    for index, existing in enumerate(list(row)):
        if existing.tag != qname("c"):
            continue
        existing_ref = existing.attrib.get("r", "")
        if existing_ref == ref:
            row.remove(existing)
            row.insert(index, cell)
            return
        if cell_column(existing_ref) > cell_column(ref):
            row.insert(index, cell)
            return
    row.append(cell)


def get_or_create_row(sheet_data: ET.Element, row_number: int) -> ET.Element:
    row_ref = str(row_number)
    for index, row in enumerate(list(sheet_data)):
        if row.tag != qname("row"):
            continue
        existing = int(row.attrib["r"])
        if existing == row_number:
            return row
        if existing > row_number:
            new_row = ET.Element(qname("row"), {"r": row_ref})
            sheet_data.insert(index, new_row)
            return new_row
    new_row = ET.Element(qname("row"), {"r": row_ref})
    sheet_data.append(new_row)
    return new_row


def find_sheet_path(xlsx: Path, sheet_name: str) -> str:
    with ZipFile(xlsx) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in relationships}
        for sheet in workbook.find("main:sheets", NS):
            if sheet.attrib["name"] == sheet_name:
                rel_id = sheet.attrib[f"{{{REL_NS}}}id"]
                return "xl/" + rel_map[rel_id].lstrip("/")
    raise ValueError(f"Sheet not found: {sheet_name}")


def build_results():
    market = pd.read_csv(
        OHLCV_CSV,
        usecols=["symbol", "date", "high"],
        dtype={"symbol": "string"},
    )
    market["symbol_norm"] = market["symbol"].str.strip().str.upper()
    market["date"] = pd.to_datetime(market["date"], errors="coerce").dt.date
    market["high"] = pd.to_numeric(market["high"], errors="coerce")
    market = market.dropna(subset=["symbol_norm", "date", "high"])
    latest_market_date = max(market["date"])
    by_symbol = {
        symbol: data.sort_values("date")[["date", "high"]].reset_index(drop=True)
        for symbol, data in market.groupby("symbol_norm", sort=False)
    }

    value_workbook = load_workbook(INPUT_XLSX, data_only=True, read_only=False)
    value_sheet = value_workbook["ROI Journal"]
    rows = {}
    calculated = 0
    missing_input = 0
    missing_market = 0

    for row in range(2, value_sheet.max_row + 1):
        stock = normalize_symbol(value_sheet.cell(row, 5).value)
        if not stock:
            continue

        identified_date = parse_workbook_date(value_sheet.cell(row, 1).value)
        entry_price = parse_number(value_sheet.cell(row, 6).value)
        result = {
            "latest_market_date": latest_market_date,
            "max_date": None,
            "max_high": None,
            "max_return": None,
            "days": None,
        }
        if identified_date is None or entry_price is None or entry_price == 0:
            missing_input += 1
        else:
            symbol_data = by_symbol.get(stock)
            if symbol_data is None:
                missing_market += 1
            else:
                window = symbol_data[
                    (symbol_data["date"] > identified_date)
                    & (symbol_data["date"] <= latest_market_date)
                ]
                if window.empty:
                    missing_market += 1
                else:
                    max_index = window["high"].idxmax()
                    max_date = window.at[max_index, "date"]
                    max_high = float(window.at[max_index, "high"])
                    result.update(
                        {
                            "max_date": max_date,
                            "max_high": max_high,
                            "max_return": (max_high - entry_price) / entry_price,
                            "days": (max_date - identified_date).days,
                        }
                    )
                    calculated += 1
        rows[row] = result

    return rows, latest_market_date, calculated, missing_input, missing_market


def patch_roi_sheet(xml_bytes: bytes, results: dict[int, dict]) -> bytes:
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        raise ValueError("ROI Journal sheetData not found")

    cols = root.find("main:cols", NS)
    if cols is None:
        cols = ET.Element(qname("cols"))
        sheet_data_index = list(root).index(sheet_data)
        root.insert(sheet_data_index, cols)
    cols.append(ET.Element(qname("col"), {"min": "42", "max": "42", "width": "14", "customWidth": "1"}))
    cols.append(ET.Element(qname("col"), {"min": "43", "max": "43", "width": "12", "customWidth": "1"}))
    cols.append(ET.Element(qname("col"), {"min": "44", "max": "44", "width": "16", "customWidth": "1"}))
    cols.append(ET.Element(qname("col"), {"min": "45", "max": "45", "width": "16", "customWidth": "1"}))
    cols.append(ET.Element(qname("col"), {"min": "46", "max": "46", "width": "16", "customWidth": "1"}))

    header_style = "7"
    date_style = "23"
    price_style = "39"
    percent_style = "68"
    integer_style = "53"

    header_row = get_or_create_row(sheet_data, 1)
    headers = [
        ("AP1", "Max_High_Date"),
        ("AQ1", "Max_High"),
        ("AR1", "Max_Market_Return"),
        ("AS1", "Max_Return_Days"),
        ("AT1", "OHLCV_Latest_Date"),
    ]
    for ref, text in headers:
        set_cell(header_row, make_inline_string_cell(ref, text, header_style))

    for row_number, result in results.items():
        row = get_or_create_row(sheet_data, row_number)
        if result["max_date"] is not None:
            set_cell(row, make_number_cell(f"AP{row_number}", excel_serial(result["max_date"]), date_style))
            set_cell(row, make_number_cell(f"AQ{row_number}", result["max_high"], price_style))
            set_cell(row, make_number_cell(f"AR{row_number}", result["max_return"], percent_style))
            set_cell(row, make_number_cell(f"AS{row_number}", result["days"], integer_style))
        set_cell(row, make_number_cell(f"AT{row_number}", excel_serial(result["latest_market_date"]), date_style))

    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def patch_notes_sheet(xml_bytes: bytes, latest_market_date: date) -> bytes:
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        return xml_bytes

    existing_rows = [
        int(row.attrib["r"])
        for row in sheet_data.findall("main:row", NS)
        if row.attrib.get("r", "").isdigit()
    ]
    note_row_number = max(existing_rows or [0]) + 1
    row = get_or_create_row(sheet_data, note_row_number)
    note = (
        "Max_Market_Return source: NSE_daily_ohlcv (1).csv. "
        f"Uses the highest high after Date_Identified through {latest_market_date.isoformat()} "
        "relative to Price_Asof."
    )
    set_cell(row, make_inline_string_cell(f"A{note_row_number}", note, "6"))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    results, latest_market_date, calculated, missing_input, missing_market = build_results()
    roi_sheet_path = find_sheet_path(INPUT_XLSX, "ROI Journal")
    notes_sheet_path = find_sheet_path(INPUT_XLSX, "Notes")

    with TemporaryDirectory() as tmp:
        tmp_output = Path(tmp) / OUTPUT_XLSX.name
        with ZipFile(INPUT_XLSX, "r") as source, ZipFile(tmp_output, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == roi_sheet_path:
                    data = patch_roi_sheet(data, results)
                elif item.filename == notes_sheet_path:
                    data = patch_notes_sheet(data, latest_market_date)
                target.writestr(item, data)
        shutil.copy2(tmp_output, OUTPUT_XLSX)

    print(f"output={OUTPUT_XLSX}")
    print(f"latest_market_date={latest_market_date}")
    print(f"calculated_rows={calculated}")
    print(f"rows_missing_input={missing_input}")
    print(f"rows_without_market_data={missing_market}")
    print("new_columns=AP:AT")


if __name__ == "__main__":
    main()
