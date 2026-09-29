from __future__ import annotations

import shutil
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook


INPUT_XLSX = Path("/Users/madhubhatt/Desktop/Weekly Buy Tracker and Technical Analysis.xlsx")
OHLCV_CSV = Path("/Users/madhubhatt/Desktop/NSE_daily_ohlcv (1).csv")
OUTPUT_DIR = Path("/Users/madhubhatt/Documents/stock_signals/outputs/between_high_low")
OUTPUT_XLSX = OUTPUT_DIR / "Weekly Buy Tracker and Technical Analysis - Between High Low.xlsx"

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"main": MAIN_NS, "rel": REL_NS}

ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)


def qname(tag: str) -> str:
    return f"{{{MAIN_NS}}}{tag}"


def parse_workbook_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%B-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def parse_number(value):
    if value is None or value == "":
        return None
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_symbol(value) -> str:
    return str(value).strip().upper() if value is not None else ""


def column_index(column_letters: str) -> int:
    total = 0
    for char in column_letters:
        total = total * 26 + ord(char.upper()) - ord("A") + 1
    return total


def cell_column(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha())
    return column_index(letters)


def make_number_cell(ref: str, value, style: str | None = None) -> ET.Element:
    cell = ET.Element(qname("c"), {"r": ref})
    if style is not None:
        cell.set("s", style)
    v = ET.SubElement(cell, qname("v"))
    if isinstance(value, float):
        v.text = f"{value:.12g}"
    else:
        v.text = str(value)
    return cell


def make_inline_string_cell(ref: str, value: str, style: str | None = None) -> ET.Element:
    cell = ET.Element(qname("c"), {"r": ref, "t": "inlineStr"})
    if style is not None:
        cell.set("s", style)
    inline = ET.SubElement(cell, qname("is"))
    text = ET.SubElement(inline, qname("t"))
    text.text = value
    return cell


def set_cell(row: ET.Element, cell: ET.Element):
    ref = cell.attrib["r"]
    for index, existing in enumerate(list(row)):
        if existing.tag != qname("c"):
            continue
        existing_ref = existing.attrib.get("r", "")
        if existing_ref == ref:
            row.remove(existing)
            row.insert(index, cell)
            return
        if cell_column(existing_ref) > cell_column(ref):
            row.insert(index, cell)
            return
    row.append(cell)


def get_or_create_row(sheet_data: ET.Element, row_number: int) -> ET.Element:
    row_ref = str(row_number)
    for index, row in enumerate(list(sheet_data)):
        if row.tag != qname("row"):
            continue
        existing = int(row.attrib["r"])
        if existing == row_number:
            return row
        if existing > row_number:
            new_row = ET.Element(qname("row"), {"r": row_ref})
            sheet_data.insert(index, new_row)
            return new_row
    new_row = ET.Element(qname("row"), {"r": row_ref})
    sheet_data.append(new_row)
    return new_row


def find_sheet_path(xlsx: Path, sheet_name: str) -> str:
    with ZipFile(xlsx) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in relationships}
        for sheet in workbook.find("main:sheets", NS):
            if sheet.attrib["name"] == sheet_name:
                rel_id = sheet.attrib[f"{{{REL_NS}}}id"]
                return "xl/" + rel_map[rel_id].lstrip("/")
    raise ValueError(f"Sheet not found: {sheet_name}")


def header_map(sheet) -> dict[str, int]:
    headers = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(1, column).value
        if isinstance(value, str) and value.strip():
            headers[value.strip().upper()] = column
    return headers


def load_source_rows() -> tuple[dict[int, dict], dict[str, int]]:
    workbook = load_workbook(INPUT_XLSX, data_only=True, read_only=False)
    sheet = workbook["ROI Journal"]
    headers = header_map(sheet)
    required = ["DATE_IDENTIFIED", "STOCKS", "BETWEEN_HIGH"]
    missing = [name for name in required if name not in headers]
    if missing:
        workbook.close()
        raise ValueError(f"Missing ROI Journal headers: {', '.join(missing)}")

    rows = {}
    try:
        for row_number in range(2, sheet.max_row + 1):
            symbol = normalize_symbol(sheet.cell(row_number, headers["STOCKS"]).value)
            if not symbol:
                continue
            rows[row_number] = {
                "symbol": symbol,
                "date_identified": parse_workbook_date(sheet.cell(row_number, headers["DATE_IDENTIFIED"]).value),
            }
    finally:
        workbook.close()
    return rows, headers


def compute_between_values(rows: dict[int, dict]) -> tuple[dict[int, dict], date]:
    market = pd.read_csv(
        OHLCV_CSV,
        usecols=["symbol", "date", "high", "low"],
        dtype={"symbol": "string"},
    )
    market["symbol_norm"] = market["symbol"].str.strip().str.upper()
    market["date"] = pd.to_datetime(market["date"], errors="coerce").dt.date
    market["high"] = pd.to_numeric(market["high"], errors="coerce")
    market["low"] = pd.to_numeric(market["low"], errors="coerce")
    market = market.dropna(subset=["symbol_norm", "date", "high", "low"])
    latest_market_date = max(market["date"])
    by_symbol = {
        symbol: data.sort_values("date")[["date", "high", "low"]].reset_index(drop=True)
        for symbol, data in market.groupby("symbol_norm", sort=False)
    }

    for result in rows.values():
        result["between_high"] = None
        result["between_low"] = None
        identified = result["date_identified"]
        symbol_data = by_symbol.get(result["symbol"])
        if identified is None or symbol_data is None:
            continue
        window = symbol_data[
            (symbol_data["date"] > identified)
            & (symbol_data["date"] <= latest_market_date)
        ]
        if window.empty:
            continue
        result["between_high"] = float(window["high"].max())
        result["between_low"] = float(window["low"].min())
    return rows, latest_market_date


def patch_roi_sheet(xml_bytes: bytes, results: dict[int, dict], between_high_col: int, between_low_col: int) -> bytes:
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        raise ValueError("ROI Journal sheetData not found")

    cols = root.find("main:cols", NS)
    if cols is None:
        cols = ET.Element(qname("cols"))
        sheet_data_index = list(root).index(sheet_data)
        root.insert(sheet_data_index, cols)
    cols.append(
        ET.Element(
            qname("col"),
            {"min": str(between_low_col), "max": str(between_low_col), "width": "14", "customWidth": "1"},
        )
    )

    header_style = "7"
    price_style = "39"
    high_col_letter = "U"
    low_col_letter = "V"

    header_row = get_or_create_row(sheet_data, 1)
    set_cell(header_row, make_inline_string_cell(f"{high_col_letter}1", "BETWEEN_HIGH", header_style))
    set_cell(header_row, make_inline_string_cell(f"{low_col_letter}1", "BETWEEN_LOW", header_style))

    for row_number, result in results.items():
        row = get_or_create_row(sheet_data, row_number)
        if result["between_high"] is not None:
            set_cell(row, make_number_cell(f"{high_col_letter}{row_number}", result["between_high"], price_style))
        if result["between_low"] is not None:
            set_cell(row, make_number_cell(f"{low_col_letter}{row_number}", result["between_low"], price_style))

    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def patch_notes_sheet(xml_bytes: bytes, latest_market_date: date) -> bytes:
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        return xml_bytes

    existing_rows = [
        int(row.attrib["r"])
        for row in sheet_data.findall("main:row", NS)
        if row.attrib.get("r", "").isdigit()
    ]
    note_row_number = max(existing_rows or [0]) + 1
    note = (
        "BETWEEN_HIGH and BETWEEN_LOW source: NSE_daily_ohlcv (1).csv. "
        f"Uses OHLCV high/low strictly after Date_Identified through {latest_market_date.isoformat()}."
    )
    row = get_or_create_row(sheet_data, note_row_number)
    set_cell(row, make_inline_string_cell(f"A{note_row_number}", note, "6"))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    source_rows, headers = load_source_rows()
    between_high_col = headers["BETWEEN_HIGH"]
    between_low_col = between_high_col + 1
    if between_high_col != 21 or between_low_col != 22:
        raise ValueError(
            f"Unexpected BETWEEN_HIGH position: {between_high_col}; this patch expects column U."
        )

    results, latest_market_date = compute_between_values(source_rows)
    populated = sum(
        1
        for result in results.values()
        if result["between_high"] is not None and result["between_low"] is not None
    )
    missing = len(results) - populated
    roi_sheet_path = find_sheet_path(INPUT_XLSX, "ROI Journal")
    notes_sheet_path = find_sheet_path(INPUT_XLSX, "Notes")

    with TemporaryDirectory() as tmp:
        tmp_output = Path(tmp) / OUTPUT_XLSX.name
        with ZipFile(INPUT_XLSX, "r") as source, ZipFile(tmp_output, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == roi_sheet_path:
                    data = patch_roi_sheet(data, results, between_high_col, between_low_col)
                elif item.filename == notes_sheet_path:
                    data = patch_notes_sheet(data, latest_market_date)
                target.writestr(deepcopy(item), data)
        shutil.copy2(tmp_output, OUTPUT_XLSX)

    print(f"output={OUTPUT_XLSX}")
    print(f"latest_market_date={latest_market_date}")
    print(f"roi_rows_with_stocks={len(results)}")
    print(f"rows_populated={populated}")
    print(f"rows_blank_no_tplus1_data={missing}")
    print("updated_columns=U:V")


if __name__ == "__main__":
    main()
