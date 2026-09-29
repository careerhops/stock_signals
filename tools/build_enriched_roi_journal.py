from __future__ import annotations

import csv
import re
import shutil
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook


INPUT_XLSX = Path("/Users/madhubhatt/Desktop/Weekly Buy Tracker and Technical Analysis.xlsx")
OHLCV_CSV = Path("/Users/madhubhatt/Desktop/NSE_daily_ohlcv (1).csv")
PROMOTER_CSV = Path("/Users/madhubhatt/Documents/stock_signals/data/instruments/promoter_holdings.csv")
INSTRUMENTS_CSV = Path("/Users/madhubhatt/Documents/stock_signals/data/instruments/instruments.csv")
DESKTOP_DIR = Path("/Users/madhubhatt/Desktop")
OUTPUT_DIR = Path("/Users/madhubhatt/Documents/stock_signals/outputs/roi_journal_enriched")
OUTPUT_XLSX = OUTPUT_DIR / "Weekly Buy Tracker and Technical Analysis - Enriched ROI Journal.xlsx"

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"main": MAIN_NS, "rel": REL_NS, "pkg": PKG_REL_NS}

ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)


@dataclass
class PromoterRecord:
    company_name: str
    holding_pct: float | None
    as_on_date: date | None
    as_on_text: str | None


@dataclass
class RoceRecord:
    company_name: str
    key: str
    roce: float
    report_date: date
    source_file: str


def qname(tag: str) -> str:
    return f"{{{MAIN_NS}}}{tag}"


def parse_workbook_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%B-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def parse_number(value):
    if value is None or value == "":
        return None
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_symbol(value) -> str:
    return str(value).strip().upper() if value is not None else ""


def excel_serial(dt: date) -> int:
    return (dt - date(1899, 12, 30)).days


def column_index(column_letters: str) -> int:
    total = 0
    for char in column_letters:
        total = total * 26 + ord(char.upper()) - ord("A") + 1
    return total


def cell_column(cell_ref: str) -> int:
    letters = "".join(ch for ch in cell_ref if ch.isalpha())
    return column_index(letters)


def make_number_cell(ref: str, value, style: str | None = None) -> ET.Element:
    cell = ET.Element(qname("c"), {"r": ref})
    if style is not None:
        cell.set("s", style)
    v = ET.SubElement(cell, qname("v"))
    if isinstance(value, float):
        v.text = f"{value:.12g}"
    else:
        v.text = str(value)
    return cell


def make_inline_string_cell(ref: str, value: str, style: str | None = None) -> ET.Element:
    cell = ET.Element(qname("c"), {"r": ref, "t": "inlineStr"})
    if style is not None:
        cell.set("s", style)
    inline = ET.SubElement(cell, qname("is"))
    text = ET.SubElement(inline, qname("t"))
    text.text = value
    return cell


def set_cell(row: ET.Element, cell: ET.Element):
    ref = cell.attrib["r"]
    for index, existing in enumerate(list(row)):
        if existing.tag != qname("c"):
            continue
        existing_ref = existing.attrib.get("r", "")
        if existing_ref == ref:
            row.remove(existing)
            row.insert(index, cell)
            return
        if cell_column(existing_ref) > cell_column(ref):
            row.insert(index, cell)
            return
    row.append(cell)


def remove_cells(row: ET.Element, min_col: int, max_col: int):
    for cell in list(row):
        if cell.tag == qname("c") and min_col <= cell_column(cell.attrib.get("r", "")) <= max_col:
            row.remove(cell)


def get_or_create_row(sheet_data: ET.Element, row_number: int) -> ET.Element:
    row_ref = str(row_number)
    for index, row in enumerate(list(sheet_data)):
        if row.tag != qname("row"):
            continue
        existing = int(row.attrib["r"])
        if existing == row_number:
            return row
        if existing > row_number:
            new_row = ET.Element(qname("row"), {"r": row_ref})
            sheet_data.insert(index, new_row)
            return new_row
    new_row = ET.Element(qname("row"), {"r": row_ref})
    sheet_data.append(new_row)
    return new_row


def find_sheet_path(xlsx: Path, sheet_name: str) -> str:
    with ZipFile(xlsx) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in relationships}
        for sheet in workbook.find("main:sheets", NS):
            if sheet.attrib["name"] == sheet_name:
                rel_id = sheet.attrib[f"{{{REL_NS}}}id"]
                return "xl/" + rel_map[rel_id].lstrip("/")
    raise ValueError(f"Sheet not found: {sheet_name}")


def normalize_company_name(value: str | None) -> str:
    if not value:
        return ""
    text = value.upper().replace("&", " AND ")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    stop_words = {
        "LTD",
        "LIMITED",
        "PVT",
        "PRIVATE",
        "CO",
        "COMPANY",
        "CORP",
        "CORPORATION",
        "INDIA",
        "INDIAN",
        "THE",
    }
    tokens = [token for token in text.split() if token not in stop_words]
    return " ".join(tokens)


def load_promoter_records() -> dict[str, PromoterRecord]:
    records: dict[str, PromoterRecord] = {}
    with PROMOTER_CSV.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            symbol = normalize_symbol(row.get("symbol"))
            if not symbol:
                continue
            holding_pct = parse_number(row.get("promoter_holding_pct"))
            as_on_text = (row.get("as_on_date") or "").strip() or None
            records[symbol] = PromoterRecord(
                company_name=(row.get("company_name") or "").strip(),
                holding_pct=holding_pct,
                as_on_date=parse_workbook_date(as_on_text),
                as_on_text=as_on_text,
            )
    return records


def load_instrument_names() -> dict[str, str]:
    names: dict[str, str] = {}
    with INSTRUMENTS_CSV.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            symbol = normalize_symbol(row.get("tradingsymbol"))
            name = (row.get("name") or "").strip()
            exchange = (row.get("exchange") or "").strip().upper()
            instrument_type = (row.get("instrument_type") or "").strip().upper()
            if symbol and name and exchange == "NSE" and instrument_type == "EQ":
                names.setdefault(symbol, name)
    return names


def numeric_cell(sheet, row: int, column: int) -> float | None:
    return parse_number(sheet.cell(row, column).value)


def read_roce_from_screener(path: Path) -> RoceRecord | None:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return None

    try:
        if "Data Sheet" not in workbook.sheetnames:
            return None
        sheet = workbook["Data Sheet"]
        company_name = str(sheet["B1"].value or "").strip()
        if not company_name:
            return None

        annual_columns = []
        for column in range(2, (sheet.max_column or 0) + 1):
            report_dt = parse_workbook_date(sheet.cell(16, column).value)
            if report_dt is None:
                continue
            pbt = numeric_cell(sheet, 28, column)
            interest = numeric_cell(sheet, 27, column)
            capital_parts = [numeric_cell(sheet, row, column) for row in (57, 58, 59)]
            if pbt is None or interest is None or any(value is None for value in capital_parts):
                continue
            annual_columns.append(
                {
                    "column": column,
                    "date": report_dt,
                    "pbt": pbt,
                    "interest": interest,
                    "capital_employed": sum(capital_parts),
                }
            )

        if len(annual_columns) < 2:
            return None
        annual_columns.sort(key=lambda item: item["date"])
        current = annual_columns[-1]
        previous = annual_columns[-2]
        denominator = current["capital_employed"] + previous["capital_employed"]
        if denominator == 0:
            return None

        roce = (current["pbt"] + current["interest"]) * 2 / denominator
        return RoceRecord(
            company_name=company_name,
            key=normalize_company_name(company_name),
            roce=roce,
            report_date=current["date"],
            source_file=path.name,
        )
    finally:
        workbook.close()


def load_roce_records() -> list[RoceRecord]:
    records_by_key: dict[str, RoceRecord] = {}
    for path in sorted(DESKTOP_DIR.glob("*.xlsx")):
        if path.name == INPUT_XLSX.name:
            continue
        record = read_roce_from_screener(path)
        if record is None or not record.key:
            continue
        # Duplicate Desktop exports such as "(1)" contain the same company data.
        records_by_key.setdefault(record.key, record)
    return list(records_by_key.values())


def match_roce(
    symbol: str,
    promoter_records: dict[str, PromoterRecord],
    instrument_names: dict[str, str],
    roce_records: list[RoceRecord],
    exact_roce: dict[str, RoceRecord],
) -> RoceRecord | None:
    candidate_names = []
    promoter = promoter_records.get(symbol)
    if promoter and promoter.company_name:
        candidate_names.append(promoter.company_name)
    instrument_name = instrument_names.get(symbol)
    if instrument_name:
        candidate_names.append(instrument_name)

    best_record = None
    best_score = 0.0
    for name in candidate_names:
        key = normalize_company_name(name)
        if not key:
            continue
        if key in exact_roce:
            return exact_roce[key]
        for record in roce_records:
            if key == record.key:
                return record
            if key in record.key or record.key in key:
                score = min(len(key), len(record.key)) / max(len(key), len(record.key))
                score = max(score, 0.82)
            else:
                score = SequenceMatcher(None, key, record.key).ratio()
            if score > best_score:
                best_score = score
                best_record = record

    if best_score >= 0.82:
        return best_record
    return None


def build_results():
    market = pd.read_csv(
        OHLCV_CSV,
        usecols=["symbol", "date", "high", "close"],
        dtype={"symbol": "string"},
    )
    market["symbol_norm"] = market["symbol"].str.strip().str.upper()
    market["date"] = pd.to_datetime(market["date"], errors="coerce").dt.date
    market["high"] = pd.to_numeric(market["high"], errors="coerce")
    market["close"] = pd.to_numeric(market["close"], errors="coerce")
    market = market.dropna(subset=["symbol_norm", "date", "high", "close"])
    latest_market_date = max(market["date"])
    by_symbol = {}
    for symbol, data in market.groupby("symbol_norm", sort=False):
        symbol_data = data.sort_values("date")[["date", "high", "close"]].reset_index(drop=True)
        symbol_data["dma75"] = symbol_data["close"].rolling(window=75, min_periods=75).mean()
        by_symbol[symbol] = symbol_data

    promoter_records = load_promoter_records()
    instrument_names = load_instrument_names()
    roce_records = load_roce_records()
    exact_roce = {record.key: record for record in roce_records}

    value_workbook = load_workbook(INPUT_XLSX, data_only=True, read_only=False)
    value_sheet = value_workbook["ROI Journal"]

    rows = {}
    calculated_max_return = 0
    rows_missing_input = 0
    rows_without_market_data = 0
    promoter_matches = 0
    roce_matches = 0
    dma75_matches = 0
    dma75_above = 0
    dma75_below = 0
    dma75_at = 0
    dma75_missing = 0
    distinct_symbols: set[str] = set()

    try:
        for row in range(2, value_sheet.max_row + 1):
            stock = normalize_symbol(value_sheet.cell(row, 5).value)
            if not stock:
                continue
            distinct_symbols.add(stock)

            identified_date = parse_workbook_date(value_sheet.cell(row, 1).value)
            entry_price = parse_number(value_sheet.cell(row, 6).value)
            promoter = promoter_records.get(stock)
            roce = match_roce(stock, promoter_records, instrument_names, roce_records, exact_roce)
            result = {
                "max_date": None,
                "max_return": None,
                "promoter_holding": None,
                "promoter_as_on_date": None,
                "promoter_as_on_text": None,
                "roce": None,
                "roce_report_date": None,
                "roce_source_file": None,
                "dma75_date": None,
                "dma75_close": None,
                "dma75": None,
                "dma75_status": None,
            }

            if identified_date is None or entry_price is None or entry_price == 0:
                rows_missing_input += 1
            else:
                symbol_data = by_symbol.get(stock)
                if symbol_data is None:
                    rows_without_market_data += 1
                    dma75_missing += 1
                else:
                    dma_window = symbol_data[symbol_data["date"] <= identified_date]
                    if dma_window.empty:
                        dma75_missing += 1
                    else:
                        dma_row = dma_window.iloc[-1]
                        dma_value = parse_number(dma_row["dma75"])
                        close_value = parse_number(dma_row["close"])
                        if dma_value is None or close_value is None:
                            dma75_missing += 1
                        else:
                            result["dma75_date"] = dma_row["date"]
                            result["dma75_close"] = close_value
                            result["dma75"] = dma_value
                            if close_value > dma_value:
                                result["dma75_status"] = "Above 75 DMA"
                                dma75_above += 1
                            elif close_value < dma_value:
                                result["dma75_status"] = "Below 75 DMA"
                                dma75_below += 1
                            else:
                                result["dma75_status"] = "At 75 DMA"
                                dma75_at += 1
                            dma75_matches += 1

                    window = symbol_data[
                        (symbol_data["date"] > identified_date)
                        & (symbol_data["date"] <= latest_market_date)
                    ]
                    if window.empty:
                        rows_without_market_data += 1
                    else:
                        max_index = window["high"].idxmax()
                        max_date = window.at[max_index, "date"]
                        max_high = float(window.at[max_index, "high"])
                        result["max_date"] = max_date
                        result["max_return"] = (max_high - entry_price) / entry_price
                        calculated_max_return += 1

            if promoter and promoter.holding_pct is not None:
                result["promoter_holding"] = promoter.holding_pct / 100
                result["promoter_as_on_date"] = promoter.as_on_date
                result["promoter_as_on_text"] = promoter.as_on_text
                promoter_matches += 1

            if roce:
                result["roce"] = roce.roce
                result["roce_report_date"] = roce.report_date
                result["roce_source_file"] = roce.source_file
                roce_matches += 1

            rows[row] = result
    finally:
        value_workbook.close()

    return {
        "rows": rows,
        "latest_market_date": latest_market_date,
        "calculated_max_return": calculated_max_return,
        "rows_missing_input": rows_missing_input,
        "rows_without_market_data": rows_without_market_data,
        "promoter_matches": promoter_matches,
        "roce_matches": roce_matches,
        "dma75_matches": dma75_matches,
        "dma75_above": dma75_above,
        "dma75_below": dma75_below,
        "dma75_at": dma75_at,
        "dma75_missing": dma75_missing,
        "distinct_symbols": len(distinct_symbols),
        "screener_exports": len(roce_records),
    }


def patch_roi_sheet(xml_bytes: bytes, results: dict[int, dict]) -> bytes:
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        raise ValueError("ROI Journal sheetData not found")

    cols = root.find("main:cols", NS)
    if cols is None:
        cols = ET.Element(qname("cols"))
        sheet_data_index = list(root).index(sheet_data)
        root.insert(sheet_data_index, cols)
    widths = {
        15: "28",
        16: "15",
        17: "12",
        18: "18",
        19: "18",
        20: "15",
        21: "14",
        22: "14",
        23: "18",
    }
    for column, width in widths.items():
        cols.append(
            ET.Element(
                qname("col"),
                {"min": str(column), "max": str(column), "width": width, "customWidth": "1"},
            )
        )

    header_style = "7"
    date_style = "23"
    price_style = "39"
    percent_style = "68"

    header_row = get_or_create_row(sheet_data, 1)
    remove_cells(header_row, 15, 23)
    headers = [
        ("O1", "Maximum_Return_After_Date_Identified"),
        ("P1", "Max_Return_Date"),
        ("Q1", "ROCE"),
        ("R1", "Promoter_Holdings"),
        ("S1", "Promoter_Holdings_As_On"),
        ("T1", "DMA75_As_Of_Date"),
        ("U1", "Close_As_Of_DMA_Date"),
        ("V1", "DMA75"),
        ("W1", "DMA75_Status"),
    ]
    for ref, text in headers:
        set_cell(header_row, make_inline_string_cell(ref, text, header_style))

    for row_number, result in results.items():
        row = get_or_create_row(sheet_data, row_number)
        remove_cells(row, 15, 23)
        if result["max_return"] is not None:
            set_cell(row, make_number_cell(f"O{row_number}", result["max_return"], percent_style))
        if result["max_date"] is not None:
            set_cell(row, make_number_cell(f"P{row_number}", excel_serial(result["max_date"]), date_style))
        if result["roce"] is not None:
            set_cell(row, make_number_cell(f"Q{row_number}", result["roce"], percent_style))
        if result["promoter_holding"] is not None:
            set_cell(row, make_number_cell(f"R{row_number}", result["promoter_holding"], percent_style))
        if result["promoter_as_on_date"] is not None:
            set_cell(row, make_number_cell(f"S{row_number}", excel_serial(result["promoter_as_on_date"]), date_style))
        elif result["promoter_as_on_text"]:
            set_cell(row, make_inline_string_cell(f"S{row_number}", result["promoter_as_on_text"]))
        if result["dma75_date"] is not None:
            set_cell(row, make_number_cell(f"T{row_number}", excel_serial(result["dma75_date"]), date_style))
        if result["dma75_close"] is not None:
            set_cell(row, make_number_cell(f"U{row_number}", result["dma75_close"], price_style))
        if result["dma75"] is not None:
            set_cell(row, make_number_cell(f"V{row_number}", result["dma75"], price_style))
        if result["dma75_status"]:
            set_cell(row, make_inline_string_cell(f"W{row_number}", result["dma75_status"]))

    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def patch_notes_sheet(xml_bytes: bytes, latest_market_date: date, screener_exports: int) -> bytes:
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        return xml_bytes

    existing_rows = [
        int(row.attrib["r"])
        for row in sheet_data.findall("main:row", NS)
        if row.attrib.get("r", "").isdigit()
    ]
    note_row_number = max(existing_rows or [0]) + 1
    notes = [
        (
            "Maximum_Return_After_Date_Identified source: NSE_daily_ohlcv (1).csv. "
            f"Uses the highest high after Date_Identified through {latest_market_date.isoformat()} "
            "relative to Price_Asof."
        ),
        (
            "Promoter_Holdings source: data/instruments/promoter_holdings.csv "
            "from NSE corporate-share-holdings-master."
        ),
        (
            "ROCE source: matching Screener export workbooks on Desktop. "
            "Computed as latest annual (Profit before tax + Interest) divided by average capital employed; "
            f"{screener_exports} Screener exports were available."
        ),
        (
            "DMA75 source: NSE_daily_ohlcv (1).csv. "
            "Uses the 75-trading-day moving average of close and compares it with the close "
            "on the latest available trading date on or before Date_Identified."
        ),
    ]
    for offset, note in enumerate(notes):
        row_number = note_row_number + offset
        row = get_or_create_row(sheet_data, row_number)
        set_cell(row, make_inline_string_cell(f"A{row_number}", note, "6"))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def write_output(stats: dict):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    roi_sheet_path = find_sheet_path(INPUT_XLSX, "ROI Journal")
    notes_sheet_path = find_sheet_path(INPUT_XLSX, "Notes")

    with TemporaryDirectory() as tmp:
        tmp_output = Path(tmp) / OUTPUT_XLSX.name
        with ZipFile(INPUT_XLSX, "r") as source, ZipFile(tmp_output, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == roi_sheet_path:
                    data = patch_roi_sheet(data, stats["rows"])
                elif item.filename == notes_sheet_path:
                    data = patch_notes_sheet(data, stats["latest_market_date"], stats["screener_exports"])
                target.writestr(deepcopy(item), data)
        shutil.copy2(tmp_output, OUTPUT_XLSX)


def main():
    stats = build_results()
    write_output(stats)
    print(f"output={OUTPUT_XLSX}")
    print(f"latest_market_date={stats['latest_market_date']}")
    print(f"roi_rows_updated={len(stats['rows'])}")
    print(f"distinct_symbols={stats['distinct_symbols']}")
    print(f"max_return_populated={stats['calculated_max_return']}")
    print(f"rows_missing_input={stats['rows_missing_input']}")
    print(f"rows_without_market_data={stats['rows_without_market_data']}")
    print(f"promoter_rows_populated={stats['promoter_matches']}")
    print(f"roce_rows_populated={stats['roce_matches']}")
    print(f"dma75_rows_populated={stats['dma75_matches']}")
    print(f"dma75_above={stats['dma75_above']}")
    print(f"dma75_below={stats['dma75_below']}")
    print(f"dma75_at={stats['dma75_at']}")
    print(f"dma75_missing={stats['dma75_missing']}")
    print(f"screener_exports_available={stats['screener_exports']}")
    print("new_columns=O:W")


if __name__ == "__main__":
    main()
