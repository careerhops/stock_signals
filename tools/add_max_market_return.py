from __future__ import annotations

from copy import copy
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


INPUT_XLSX = Path("/Users/madhubhatt/Desktop/Weekly Buy Tracker and Technical Analysis.xlsx")
OHLCV_CSV = Path("/Users/madhubhatt/Desktop/NSE_daily_ohlcv (1).csv")
OUTPUT_DIR = Path("/Users/madhubhatt/Documents/stock_signals/outputs/max_market_return")
OUTPUT_XLSX = OUTPUT_DIR / "Weekly Buy Tracker and Technical Analysis - Max Market Return.xlsx"


def parse_workbook_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def parse_number(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def copy_cell_style(source, target):
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.font:
        target.font = copy(source.font)
    if source.fill:
        target.fill = copy(source.fill)
    if source.border:
        target.border = copy(source.border)
    if source.alignment:
        target.alignment = copy(source.alignment)


def normalize_symbol(value):
    return str(value).strip().upper() if value is not None else ""


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    market = pd.read_csv(
        OHLCV_CSV,
        usecols=["symbol", "date", "high"],
        dtype={"symbol": "string"},
    )
    market["symbol_norm"] = market["symbol"].str.strip().str.upper()
    market["date"] = pd.to_datetime(market["date"], errors="coerce").dt.date
    market["high"] = pd.to_numeric(market["high"], errors="coerce")
    market = market.dropna(subset=["symbol_norm", "date", "high"])
    latest_market_date = max(market["date"])

    by_symbol = {
        symbol: data.sort_values("date")[["date", "high"]].reset_index(drop=True)
        for symbol, data in market.groupby("symbol_norm", sort=False)
    }

    workbook = load_workbook(INPUT_XLSX, data_only=False, keep_links=True)
    value_workbook = load_workbook(INPUT_XLSX, data_only=True, keep_links=True)
    sheet = workbook["ROI Journal"]
    value_sheet = value_workbook["ROI Journal"]

    headers = [
        "Max_High_Date",
        "Max_High",
        "Max_Market_Return",
        "Max_Return_Days",
        "OHLCV_Latest_Date",
    ]
    existing_headers = {sheet.cell(1, col).value: col for col in range(1, sheet.max_column + 1)}
    if all(header in existing_headers for header in headers):
        start_col = existing_headers[headers[0]]
    else:
        start_col = sheet.max_column + 1

    header_style_source = sheet.cell(1, 14)
    body_style_source_col = 14
    for offset, header in enumerate(headers):
        cell = sheet.cell(1, start_col + offset)
        cell.value = header
        copy_cell_style(header_style_source, cell)
        cell.alignment = copy(header_style_source.alignment)

    calculated = 0
    no_market_rows = 0
    no_input_rows = 0

    for row in range(2, sheet.max_row + 1):
        stock = normalize_symbol(value_sheet.cell(row, 5).value)
        identified_date = parse_workbook_date(value_sheet.cell(row, 1).value)
        entry_price = parse_number(value_sheet.cell(row, 6).value)

        values = [None, None, None, None, None]
        if not stock:
            pass
        elif identified_date is None or entry_price is None or entry_price == 0:
            no_input_rows += 1
            values[-1] = latest_market_date
        else:
            symbol_data = by_symbol.get(stock)
            if symbol_data is None:
                no_market_rows += 1
                values[-1] = latest_market_date
            else:
                window = symbol_data[
                    (symbol_data["date"] >= identified_date)
                    & (symbol_data["date"] <= latest_market_date)
                ]
                if window.empty:
                    no_market_rows += 1
                    values[-1] = latest_market_date
                else:
                    max_index = window["high"].idxmax()
                    max_date = window.at[max_index, "date"]
                    max_high = float(window.at[max_index, "high"])
                    max_return = (max_high - entry_price) / entry_price
                    values = [
                        max_date,
                        max_high,
                        max_return,
                        (max_date - identified_date).days,
                        latest_market_date,
                    ]
                    calculated += 1

        for offset, value in enumerate(values):
            target = sheet.cell(row, start_col + offset)
            source = sheet.cell(row, body_style_source_col)
            copy_cell_style(source, target)
            target.value = value

    sheet.column_dimensions[get_column_letter(start_col)].width = 14
    sheet.column_dimensions[get_column_letter(start_col + 1)].width = 12
    sheet.column_dimensions[get_column_letter(start_col + 2)].width = 16
    sheet.column_dimensions[get_column_letter(start_col + 3)].width = 16
    sheet.column_dimensions[get_column_letter(start_col + 4)].width = 16

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, start_col).number_format = "dd/mm/yyyy"
        sheet.cell(row, start_col + 1).number_format = "0.00"
        sheet.cell(row, start_col + 2).number_format = "0.00%"
        sheet.cell(row, start_col + 3).number_format = "0"
        sheet.cell(row, start_col + 4).number_format = "dd/mm/yyyy"

    notes = workbook["Notes"] if "Notes" in workbook.sheetnames else None
    if notes is not None:
        note_row = 1
        while notes.cell(note_row, 1).value:
            note_row += 1
        notes.cell(note_row, 1).value = (
            "Max_Market_Return source: NSE_daily_ohlcv (1).csv. "
            f"Uses the highest high from Date_Identified through {latest_market_date.isoformat()} "
            "relative to Price_Asof."
        )

    workbook.save(OUTPUT_XLSX)
    print(f"output={OUTPUT_XLSX}")
    print(f"latest_market_date={latest_market_date}")
    print(f"calculated_rows={calculated}")
    print(f"rows_missing_input={no_input_rows}")
    print(f"rows_without_market_data={no_market_rows}")
    print(f"new_columns={get_column_letter(start_col)}:{get_column_letter(start_col + len(headers) - 1)}")


if __name__ == "__main__":
    main()
